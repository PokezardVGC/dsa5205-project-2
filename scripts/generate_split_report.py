from docx import Document
from docx.shared import Inches
import os

def build_report():
    doc = Document()
    doc.add_heading('Group 28: NVDA Buy the Dip Strategy Report (Revised & Separated)', 0)

    with open('report_text.txt', 'r') as f:
        text = f.read()

    doc.add_heading('Introduction', level=1)
    doc.add_paragraph("This report evaluates a mean-reversion 'Buy the Dip' strategy applied to NVDA. Based on the latest Quantitative and Machine Learning feedback, we have separated our visual and performance analysis into distinct, highly readable charts.")

    doc.add_heading('Exploratory Data Analysis (EDA)', level=1)
    doc.add_paragraph("The EDA has been separated to clearly demonstrate the rationale behind a mean-reversion approach for NVDA.")
    
    if os.path.exists('eda_1_price.png'):
        doc.add_picture('eda_1_price.png', width=Inches(6.0))
        doc.add_paragraph("Figure 1: Log-scale price chart of NVDA with 50-day and 200-day Simple Moving Averages. This context is essential for understanding the overall asset trend against its short-term and long-term momentum.")

    if os.path.exists('eda_2_zscore.png'):
        doc.add_picture('eda_2_zscore.png', width=Inches(6.0))
        doc.add_paragraph("Figure 2: Rolling 50-day Z-score. We define a 'dip' mathematically as instances where the Z-score crosses below the -2 standard deviation threshold, represented by the red shaded regions.")

    if os.path.exists('eda_3_hurst.png'):
        doc.add_picture('eda_3_hurst.png', width=Inches(6.0))
        doc.add_paragraph("Figure 3: Rolling Hurst Exponent. A Hurst exponent below 0.5 (shaded green) statistically confirms periods where NVDA exhibits mean-reverting behavior, justifying the underlying strategy.")

    if os.path.exists('eda_4_drawdown.png'):
        doc.add_picture('eda_4_drawdown.png', width=Inches(6.0))
        doc.add_paragraph("Figure 4: Underwater drawdown chart. This highlights the severity and duration of peak-to-trough declines for NVDA, with psychological thresholds marked at -10% and -20%.")

    if os.path.exists('eda_5_corr.png'):
        doc.add_picture('eda_5_corr.png', width=Inches(6.0))
        doc.add_paragraph("Figure 5: 60-day rolling correlation with SPY. By tracking this, we can assess whether NVDA's dips are idiosyncratic (asset-specific) or systematic (market-wide).")

    if os.path.exists('eda_6_vol.png'):
        doc.add_picture('eda_6_vol.png', width=Inches(6.0))
        doc.add_paragraph("Figure 6: Rolling annualized volatility. Dips frequently cluster with volatility spikes, an important factor for risk-management.")

    if os.path.exists('eda_dist.png'):
        doc.add_picture('eda_dist.png', width=Inches(6.0))
        doc.add_paragraph("Figure 7: Kernel Density Estimation (KDE) comparing the daily return distributions of NVDA and SPY. The pronounced fat tails and leptokurtosis in NVDA underscore the extreme risks of trying to catch a falling knife.")

    doc.add_heading('Out-of-Sample Performance Evaluation', level=1)
    doc.add_paragraph("The strategy's performance on out-of-sample data is broken down into the following key metrics:")
    
    if os.path.exists('perf_1_cumulative.png'):
        doc.add_picture('perf_1_cumulative.png', width=Inches(6.0))
        doc.add_paragraph("Figure 8: Out-of-sample cumulative returns (log scale). This clearly separates the growth trajectory of the strategy versus pure buy-and-hold benchmarks.")

    if os.path.exists('perf_2_underwater.png'):
        doc.add_picture('perf_2_underwater.png', width=Inches(6.0))
        doc.add_paragraph("Figure 9: Strategy vs. Benchmark underwater profile. The red area plots exactly when and how deeply the strategy experienced capital lock-up or loss from all-time highs.")

    if os.path.exists('perf_3_sharpe.png'):
        doc.add_picture('perf_3_sharpe.png', width=Inches(6.0))
        doc.add_paragraph("Figure 10: Rolling 6-month Sharpe Ratio. Rather than a single static number, this demonstrates the temporal stability and decay of the strategy's risk-adjusted edge.")

    if os.path.exists('perf_4_beta.png'):
        doc.add_picture('perf_4_beta.png', width=Inches(6.0))
        doc.add_paragraph("Figure 11: Rolling 6-month Beta to SPY. This ensures the strategy does not inadvertently take on massive market exposure during periods of high turbulence.")

    doc.add_heading('Conclusion & Original Text Excerpts', level=1)
    doc.add_paragraph("Below are selected excerpts from the original extracted text for continuity:")
    doc.add_paragraph(text[:3000] + "\n...[truncated for brevity]...")

    doc.save('Group_28_Report_Final.docx')
    print("Report saved to Group_28_Report_Final.docx")

if __name__ == '__main__':
    build_report()
