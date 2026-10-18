from docx import Document
from docx.shared import Inches

def build_report():
    doc = Document()
    doc.add_heading('Group 28: NVDA Buy the Dip Strategy Report (Revised)', 0)

    with open('report_text.txt', 'r') as f:
        text = f.read()

    # Add introduction
    doc.add_heading('Introduction', level=1)
    doc.add_paragraph("This report evaluates a mean-reversion 'Buy the Dip' strategy applied to NVDA. Based on latest ML and Quantitative evaluations, we have updated our visual and performance analysis.")

    doc.add_heading('Exploratory Data Analysis (EDA)', level=1)
    doc.add_paragraph("Based on feedback from Quantitative and ML analysis, we improved the EDA to show non-stationarity, tail-risks, and drawdown duration:")
    doc.add_picture('eda_price.png', width=Inches(6.0))
    doc.add_paragraph("Figure 1: Log-scale price chart with 50/200 day SMAs, Volatility Regime mapping, rolling Z-score (dip threshold), and rolling Hurst Exponent.")
    
    doc.add_picture('eda_dist.png', width=Inches(6.0))
    doc.add_paragraph("Figure 2: Kernel Density Estimation showing heavy tails and leptokurtosis in NVDA compared to SPY.")

    doc.add_picture('eda_drawdown.png', width=Inches(6.0))
    doc.add_paragraph("Figure 3: Underwater drawdown chart, rolling correlation with SPY, and rolling Volatility clustering.")

    doc.add_heading('Performance Evaluation', level=1)
    doc.add_paragraph("Our out-of-sample ML model evaluation is structured into three rigorous Tear-Sheet panels:")
    
    doc.add_picture('perf_hero.png', width=Inches(6.0))
    doc.add_paragraph("Figure 4: Hero Chart - Cumulative Returns (Log Scale) proving alpha over Buy & Hold, and the synchronized Underwater Drawdown Profile.")

    doc.add_picture('perf_rolling_risk.png', width=Inches(6.0))
    doc.add_paragraph("Figure 5: Rolling 6-Month Sharpe Ratio and Beta. This demonstrates the consistency of the risk-adjusted edge and market exposure dynamically.")

    doc.add_picture('perf_trade_dist.png', width=Inches(6.0))
    doc.add_paragraph("Figure 6: Return distribution exclusively on active trading days, detailing Win Rate and Mean Return asymmetry.")

    doc.add_heading('Conclusion & Original Text Excerpts', level=1)
    doc.add_paragraph("Below are selected excerpts from the original extracted text for continuity:")
    # add first 3000 chars of original text as reference
    doc.add_paragraph(text[:3000] + "\n...[truncated for brevity]...")

    doc.save('Group_28_Report_Final.docx')
    print("Report saved to Group_28_Report_Final.docx")

if __name__ == '__main__':
    build_report()
